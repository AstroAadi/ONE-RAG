"""
Universal Excel File Extractor
Supports: .xls, .xlsx, .xlsm, .csv
"""

from pathlib import Path
from typing import Tuple, Dict, List
import logging
import zipfile
import csv

try:
    from ..base_extractor import BaseExtractor
except ImportError:
    from base_extractor import BaseExtractor

logger = logging.getLogger(__name__)


class XLSExtractor(BaseExtractor):
    """
    Universal Excel Extractor
    Supports:
        - .xls  (legacy binary)
        - .xlsx (OpenXML)
        - .xlsm (macro-enabled OpenXML)
        - .csv
    """

    def _get_supported_formats(self) -> List[str]:
        return ['xls', 'xlsx', 'xlsm', 'csv']

    # ================= VALIDATION ================= #

    def validate_file(self, file_path: str) -> Tuple[bool, str]:
        path = Path(file_path)

        if not path.exists():
            return False, "File does not exist"

        ext = path.suffix.lower().lstrip('.')

        if ext not in self.supported_formats:
            return False, f"Unsupported file format: .{ext}"

        # Validate legacy XLS (BIFF format)
        if ext == "xls":
            try:
                with open(file_path, 'rb') as f:
                    header = f.read(8)
                    if not (header[:4] == b'\xd0\xcf\x11\xe0' or header[:2] == b'\x09\x08'):
                        return False, "Not a valid XLS file"
            except Exception as e:
                return False, f"Could not read file: {e}"

        # Validate XLSX / XLSM (must be ZIP archive)
        elif ext in ["xlsx", "xlsm"]:
            if not zipfile.is_zipfile(file_path):
                return False, "Not a valid XLSX/XLSM file"

        # CSV needs no special binary validation

        return True, f"Valid {ext.upper()} file"

    # ================= EXTRACTION ================= #

    def extract(self, file_path: str) -> Tuple[str, bool, Dict]:
        try:
            is_valid, msg = self.validate_file(file_path)
            if not is_valid:
                return "", False, {'error': msg}

            logger.info(f"Extracting spreadsheet: {file_path}")

            ext = Path(file_path).suffix.lower().lstrip('.')
            stem = Path(file_path).stem

            # Route based on extension
            if ext == "xls":
                extraction_result = self._extract_xls(file_path)
                method_used = "xlrd"

            elif ext in ["xlsx", "xlsm"]:
                extraction_result = self._extract_xlsx(file_path)
                method_used = "openpyxl"

            elif ext == "csv":
                extraction_result = self._extract_csv(file_path)
                method_used = "csv"

            else:
                return "", False, {'error': f"Unsupported format: {ext}"}

            sheets = extraction_result['sheets']

            if not sheets:
                return "", False, {'error': 'No sheets found'}

            markdown_content = self._build_markdown(stem, sheets)

            metadata = self.get_metadata(file_path)
            metadata.update({
                'sheet_count': len(sheets),
                'sheet_names': [s['name'] for s in sheets],
                'total_rows': sum(s.get('rows', 0) for s in sheets),
                'total_cols': sum(s.get('cols', 0) for s in sheets),
                'extraction_method': method_used
            })

            logger.info(f"✓ Spreadsheet extraction complete: {len(sheets)} sheets")

            return markdown_content, True, metadata

        except Exception as e:
            logger.error(f"Spreadsheet extraction error: {str(e)}")
            import traceback
            traceback.print_exc()
            return "", False, {'error': str(e)}

    # ================= XLS (LEGACY) ================= #

    def _extract_xls(self, file_path: str) -> Dict:
        import xlrd

        workbook = xlrd.open_workbook(file_path)
        sheets = []

        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)

            data = []
            for row_idx in range(sheet.nrows):
                row = []
                for col_idx in range(sheet.ncols):
                    cell = sheet.cell(row_idx, col_idx)
                    row.append(str(cell.value) if cell.value is not None else "")
                data.append(row)

            sheets.append({
                'name': sheet_name,
                'data': data,
                'rows': sheet.nrows,
                'cols': sheet.ncols
            })

        return {'sheets': sheets}

    # ================= XLSX / XLSM ================= #

    def _extract_xlsx(self, file_path: str) -> Dict:
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, data_only=True)
        sheets = []

        for sheet in workbook.worksheets:
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append([
                    str(cell) if cell is not None else ""
                    for cell in row
                ])

            sheets.append({
                'name': sheet.title,
                'data': data,
                'rows': len(data),
                'cols': len(data[0]) if data else 0
            })

        return {'sheets': sheets}

    # ================= CSV ================= #

    def _extract_csv(self, file_path: str) -> Dict:
        sheets = []
        data = []

        with open(file_path, newline='', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                data.append([str(cell) for cell in row])

        sheets.append({
            'name': 'Sheet1',
            'data': data,
            'rows': len(data),
            'cols': len(data[0]) if data else 0
        })

        return {'sheets': sheets}

    # ================= MARKDOWN BUILDER ================= #

    def _build_markdown(self, stem: str, sheets: List[Dict]) -> str:
        lines = []

        lines.append(f"# Spreadsheet: {stem}")
        lines.append(f"**Total Sheets**: {len(sheets)}")
        lines.append("")
        lines.append("---\n")

        for sheet_idx, sheet in enumerate(sheets, 1):
            lines.append(f"## Sheet {sheet_idx}: {sheet['name']}\n")
            lines.append(f"**Dimensions**: {sheet['rows']} rows × {sheet['cols']} columns\n")

            data = sheet['data']
            if not data:
                lines.append("*(Empty sheet)*\n")
                continue

            preview_data = data[:100]

            header = preview_data[0]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join(["---"] * len(header)) + " |")

            for row in preview_data[1:]:
                while len(row) < len(header):
                    row.append("")
                lines.append("| " + " | ".join(row[:len(header)]) + " |")

            if len(data) > 100:
                lines.append(f"\n*(... {len(data) - 100} more rows not shown)*")

            lines.append("")

        return "\n".join(lines)
